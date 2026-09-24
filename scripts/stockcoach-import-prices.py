#!/usr/bin/env python3
"""
Importe les prix de vente saisis dans l'export Excel Stockcoach, en ne changeant
que ce qui est nouveau par rapport a ce qui est deja publie.

Lit l'onglet "Stockcoach" : colonnes "ID source" et "Prix de vente (EUR)", et
compare avec stockcoach-prices.json (les prix actuellement publies) :
  - voiture avec un prix, pas encore publiee  -> ajoutee
  - voiture deja publiee, prix different      -> prix mis a jour
  - voiture deja publiee, meme prix           -> rien ne change
  - voiture deja publiee, sans prix / absente -> conservee telle quelle
Une voiture n'est jamais retiree par l'import : elle sort du site quand
Stockcoach la passe "Retire" (ou sur demande explicite).

Usage :
  python3 scripts/stockcoach-import-prices.py "fichier.xlsx"               montre les changements
  python3 scripts/stockcoach-import-prices.py "fichier.xlsx" --appliquer   les enregistre et publie
"""
import json
import subprocess
import sys

import openpyxl

PRICES = "stockcoach-prices.json"
CATALOG = "stockcoach-vehicles.json"

args = [a for a in sys.argv[1:] if a != "--appliquer"]
apply = "--appliquer" in sys.argv
if len(args) != 1:
    sys.exit(__doc__)

ws = openpyxl.load_workbook(args[0], data_only=True)["Stockcoach"]
header = [c.value for c in ws[1]]
id_col = header.index("ID source")
price_col = header.index("Prix de vente (EUR)")

incoming = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    source_id, price = row[id_col], row[price_col]
    if source_id in (None, "") or price in (None, ""):
        continue
    if not isinstance(price, (int, float)) or price <= 0:
        sys.exit(f"Prix invalide pour l'ID {source_id} : {price!r}")
    key = str(int(source_id))
    if key in incoming:
        sys.exit(f"ID {key} present deux fois dans le fichier")
    incoming[key] = round(price)

try:
    with open(PRICES) as f:
        current = json.load(f)
except FileNotFoundError:
    current = {}

with open(CATALOG) as f:
    catalog = {str(v["_sourceId"]): v for v in json.load(f)}


def label(sid):
    v = catalog.get(sid)
    return f"{v['make']} {v['model']} (ID {sid})" if v else f"ID {sid} (inconnu du catalogue)"


added = {k: p for k, p in incoming.items() if k not in current}
changed = {k: p for k, p in incoming.items() if k in current and current[k] != p}
same = [k for k, p in incoming.items() if current.get(k) == p]
kept = [k for k in current if k not in incoming]


def fmt(n):
    return f"{n:,}".replace(",", " ") + " EUR"


print(f"Deja publiees : {len(current)} voiture(s)\n")
print(f"Nouvelles a publier : {len(added)}")
for k, p in added.items():
    print(f"  + {label(k)} : {fmt(p)}")
print(f"Prix modifies : {len(changed)}")
for k, p in changed.items():
    print(f"  ~ {label(k)} : {fmt(current[k])} -> {fmt(p)}")
print(f"Inchangees : {len(same)}")
print(f"Deja publiees, sans prix dans ce fichier (conservees) : {len(kept)}")

if not added and not changed:
    print("\nRien de nouveau : le site reste tel quel.")
elif not apply:
    print("\nAucune modification faite. Relancer avec --appliquer pour publier ces changements.")
else:
    with open(PRICES, "w") as f:
        json.dump({**current, **incoming}, f, indent=2)
    print(f"\n{len(added)} ajout(s) et {len(changed)} changement(s) de prix enregistres.")
    subprocess.run(["node", "scripts/stockcoach-publish.js"], check=True)
